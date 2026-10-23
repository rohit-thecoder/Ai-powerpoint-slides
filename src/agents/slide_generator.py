
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"


class SlideGenerator:
    """Generate PowerPoint presentations from a slide plan."""

    def __init__(self, output_directory=None):
        self.output_directory = Path(
            output_directory or DEFAULT_OUTPUT_DIR
        )
        self.output_directory.mkdir(parents=True, exist_ok=True)

        self.navy = RGBColor(20, 35, 65)
        self.blue = RGBColor(45, 105, 210)
        self.white = RGBColor(255, 255, 255)
        self.dark_text = RGBColor(40, 48, 62)
        self.muted_text = RGBColor(105, 115, 130)
        self.background = RGBColor(247, 249, 252)
        self.light_blue = RGBColor(225, 235, 252)

    def _add_text(
        self, slide, text, left, top, width, height,
        font_size, color, bold=False,
    ):
        """Add a formatted text box to a slide."""

        shape = slide.shapes.add_textbox(
            Inches(left), Inches(top),
            Inches(width), Inches(height),
        )

        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.03)
        frame.margin_right = Inches(0.03)
        frame.margin_top = Inches(0.02)
        frame.margin_bottom = Inches(0.02)

        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = str(text)
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color

        return shape

    def _set_background(self, slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def _add_footer(self, slide, number, total):
        self._add_text(
            slide, "AI POWERPOINT GENERATOR",
            0.7, 7.05, 5.5, 0.25, 9,
            self.muted_text, bold=True,
        )

        self._add_text(
            slide, f"{number} / {total}",
            11.5, 7.05, 1.1, 0.25, 10,
            self.muted_text,
        )

    def _add_title_slide(self, presentation, title):
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )
        self._set_background(slide, self.navy)

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.85), Inches(2.1),
            Inches(0.12), Inches(2.0),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = self.blue
        accent.line.fill.background()

        self._add_text(
            slide, "AI-GENERATED PRESENTATION",
            1.25, 2.1, 10.8, 0.45, 13,
            self.light_blue, bold=True,
        )

        self._add_text(
            slide, title,
            1.25, 2.7, 10.8, 1.65, 32,
            self.white, bold=True,
        )

        self._add_text(
            slide, "Generated from your source document",
            1.25, 4.55, 10.8, 0.5, 16,
            self.light_blue,
        )

        return slide

    def _add_content_slide(
        self, presentation, title, bullets, number, total
    ):
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )
        self._set_background(slide, self.background)

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            Inches(13.333), Inches(0.12),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = self.blue
        accent.line.fill.background()

        self._add_text(
            slide, title,
            0.7, 0.5, 11.9, 0.85, 25,
            self.navy, bold=True,
        )

        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.7), Inches(1.45),
            Inches(11.9), Inches(0.025),
        )
        divider.fill.solid()
        divider.fill.fore_color.rgb = self.light_blue
        divider.line.fill.background()

        # Divide the usable vertical space between bullet points.
        count = len(bullets)
        top = 1.75
        available_height = 4.9
        item_height = available_height / count

        for index, bullet in enumerate(bullets):
            y = top + index * item_height

            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.83), Inches(y + 0.12),
                Inches(0.12), Inches(0.12),
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = self.blue
            dot.line.fill.background()

            self._add_text(
                slide, bullet,
                1.15, y, 11.0, item_height - 0.04,
                17, self.dark_text,
            )

        self._add_footer(slide, number, total)
        return slide

    def create_slide_deck(
        self,
        title,
        slides,
        filename="generated_presentation.pptx",
    ):
        """
        Expected slides:
        [
            {"title": "Slide heading", "bullets": ["Point 1", "Point 2"]}
        ]
        """

        if not isinstance(title, str) or not title.strip():
            raise ValueError("Presentation title cannot be empty.")

        if not isinstance(slides, list) or not slides:
            raise ValueError("At least one content slide is required.")

        # Validate slide content before creating the presentation.
        validated_slides = []

        for item in slides:
            if not isinstance(item, dict):
                continue

            slide_title = item.get("title")
            bullets = item.get("bullets")

            if not isinstance(slide_title, str):
                continue

            if not isinstance(bullets, list):
                continue

            bullets = [
                bullet.strip()
                for bullet in bullets
                if isinstance(bullet, str) and bullet.strip()
            ]

            if bullets:
                validated_slides.append({
                    "title": slide_title.strip(),
                    "bullets": bullets,
                })

        if not validated_slides:
            raise ValueError("No valid content slides were provided.")

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        total = len(validated_slides) + 1

        self._add_title_slide(presentation, title)

        for index, item in enumerate(validated_slides, start=2):
            self._add_content_slide(
                presentation=presentation,
                title=item["title"],
                bullets=item["bullets"],
                number=index,
                total=total,
            )

        # Use only the filename, not an arbitrary filesystem path.
        safe_filename = Path(filename).name

        if not safe_filename.lower().endswith(".pptx"):
            safe_filename += ".pptx"

        output_path = self.output_directory / safe_filename
        presentation.save(str(output_path))

        print(f"Presentation saved: {output_path}")
        print(f"Total slides: {total}")

        return output_path
